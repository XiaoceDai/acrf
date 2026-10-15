#!/usr/bin/env python3
"""Build the Frontiers DOCX from the final Markdown.

Display equations and the indexed symbols s_p and M_p are written as Word
math (OMML). Code, donor IDs, and resource names keep their underscores.
Table 1 rows do not split across pages, and the header row repeats.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
TOKEN = re.compile(r"score_dp\^LA|score_dp\^RA|aCAIrel_dp|aCAI_dp|Δ_dp|s_p|M_p")
DISPLAY = {
    "Δ_dp = score_dp^LA − score_dp^RA.": (
        "Δ_dp",
        " = ",
        "score_dp^LA",
        " − ",
        "score_dp^RA",
        ".",
    ),
    "aCAI_dp = s_p · Δ_dp,": (
        "aCAI_dp",
        " = ",
        "s_p",
        " · ",
        "Δ_dp",
        ",",
    ),
    "aCAIrel_dp = aCAI_dp / M_p.": (
        "aCAIrel_dp",
        " = ",
        "aCAI_dp",
        " / ",
        "M_p",
        ".",
    ),
}


def escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def m_run(text: str) -> str:
    return (
        "<m:r><m:rPr><m:sty m:val=\"p\"/></m:rPr>"
        "<w:rPr><w:rFonts w:ascii=\"Times New Roman\" w:hAnsi=\"Times New Roman\"/>"
        "<w:sz w:val=\"24\"/></w:rPr>"
        f"<m:t xml:space=\"preserve\">{escape(text)}</m:t></m:r>"
    )


def m_sub(base: str, sub: str) -> str:
    return f"<m:sSub><m:e>{m_run(base)}</m:e><m:sub>{m_run(sub)}</m:sub></m:sSub>"


def m_subsup(base: str, sub: str, sup: str) -> str:
    return (
        "<m:sSubSup>"
        f"<m:e>{m_run(base)}</m:e>"
        f"<m:sub>{m_run(sub)}</m:sub>"
        f"<m:sup>{m_run(sup)}</m:sup>"
        "</m:sSubSup>"
    )


def token_xml(token: str) -> str:
    return {
        "Δ_dp": m_sub("Δ", "dp"),
        "s_p": m_sub("s", "p"),
        "M_p": m_sub("M", "p"),
        "aCAI_dp": m_sub("aCAI", "dp"),
        "aCAIrel_dp": m_sub("aCAIrel", "dp"),
        "score_dp^LA": m_subsup("score", "dp", "LA"),
        "score_dp^RA": m_subsup("score", "dp", "RA"),
    }[token]


def math_element(parts: tuple[str, ...] | list[str], display: bool = False):
    inner = []
    for part in parts:
        inner.append(token_xml(part) if part in {
            "Δ_dp", "s_p", "M_p", "aCAI_dp", "aCAIrel_dp", "score_dp^LA", "score_dp^RA"
        } else m_run(part))
    body = "".join(inner)
    if display:
        xml = (
            f'<m:oMathPara xmlns:m="{M_NS}" xmlns:w="{W_NS}">'
            f"<m:oMath>{body}</m:oMath></m:oMathPara>"
        )
    else:
        xml = f'<m:oMath xmlns:m="{M_NS}" xmlns:w="{W_NS}">{body}</m:oMath>'
    return parse_xml(xml)


def set_run_font(run, name: str, size: int, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_text_run(paragraph, text: str, bold: bool = False, mono: bool = False) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, "Courier New" if mono else "Times New Roman", 10 if mono else 12, bold)


def add_plain_or_math(paragraph, text: str, bold: bool = False) -> None:
    pos = 0
    for match in TOKEN.finditer(text):
        if match.start() > pos:
            add_text_run(paragraph, text[pos:match.start()], bold=bold)
        paragraph._p.append(math_element((match.group(0),)))
        pos = match.end()
    if pos < len(text):
        add_text_run(paragraph, text[pos:], bold=bold)


def add_mixed(paragraph, text: str, bold: bool = False) -> None:
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            add_text_run(paragraph, piece[1:-1], bold=bold, mono=True)
        else:
            add_plain_or_math(paragraph, piece, bold=bold)


def add_markdown_runs(paragraph, text: str, mono: bool = False) -> None:
    if mono:
        add_text_run(paragraph, text, mono=True)
        return
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_mixed(paragraph, part[2:-2], bold=True)
        else:
            add_mixed(paragraph, part, bold=False)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_header(row) -> None:
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        tc_pr.append(shd)


def build(markdown_path: Path, docx_path: Path) -> None:
    lines = markdown_path.read_text().splitlines()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    section = doc.sections[0]
    section.page_width = Twips(12240)
    section.page_height = Twips(15840)
    section.top_margin = Twips(1440)
    section.bottom_margin = Twips(1440)
    section.left_margin = Twips(1800)
    section.right_margin = Twips(1800)

    def add_body(text: str, mono: bool = False):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        add_markdown_runs(paragraph, text, mono=mono)
        return paragraph

    i = 0
    in_code = False
    code_buf: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                add_body("\n".join(code_buf), mono=True)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line in DISPLAY:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph._p.append(math_element(DISPLAY[line], display=True))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") and c for c in cells):
                    rows.append(cells)
                i += 1
            cols = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.autofit = True
            n_resources = len(rows) - 1
            blob = "\n".join(" | ".join(row) for row in rows)
            if n_resources < 10 or "GSE261363" in blob:
                raise SystemExit(f"Table 1 has {n_resources} resources or still lists GSE261363")
            for r_index, row in enumerate(rows):
                for c_index in range(cols):
                    cell = table.cell(r_index, c_index)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    add_markdown_runs(paragraph, row[c_index] if c_index < len(row) else "")
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        if r_index == 0:
                            run.bold = True
                prevent_row_split(table.rows[r_index])
                if r_index == 0:
                    mark_header_row(table.rows[r_index])
                    shade_header(table.rows[r_index])
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            if line.strip() == "## Table 1":
                i += 1
                continue
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_markdown_runs(paragraph, line[2:].strip())
        elif line.strip():
            paragraph = add_body(line)
            if line.startswith("**Table 1."):
                paragraph.paragraph_format.page_break_before = True
                paragraph.paragraph_format.keep_with_next = True
            if line.startswith("**Figure "):
                paragraph.paragraph_format.keep_together = True
        i += 1

    section = doc.sections[0]
    sect_pr = section._sectPr
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    line_numbers.set(qn("w:distance"), "360")
    sect_pr.append(line_numbers)
    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = 1
    run = footer_paragraph.add_run()
    set_run_font(run, "Times New Roman", 10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


if __name__ == "__main__":
    import sys

    root = Path("/home/bio/Project_P02")
    markdown = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        root / "revision/acrf_submission_reviewfix/frontiers_bioinformatics/Frontiers_acrf_Technology_and_Code.md"
    )
    docx = Path(sys.argv[2]) if len(sys.argv) > 2 else (
        root / "revision/acrf_submission_reviewfix/frontiers_bioinformatics/Frontiers_acrf_Technology_and_Code.docx"
    )
    build(markdown, docx)
