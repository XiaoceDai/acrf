#!/usr/bin/env python3
"""Migrate reader-visible NCP_v3 wording while preserving internal identifiers."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
APPROVED_TITLE = (
    "Definition and freezing of the healthy paired-chamber NCP coordinate system"
)
MIGRATION_NOTE_FRAGMENT = "historical ncp_v3 identifier"

# Slashes, dots, hyphens, and underscores identify paths/object IDs that must
# retain the historical token. Standalone publication terminology is migrated.
STANDALONE_NCP_V3 = re.compile(
    r"(?<![A-Za-z0-9_./-])NCP(?:_| )v3(?![A-Za-z0-9_./-])",
    flags=re.IGNORECASE,
)

SPECIAL_PREFIX_REPLACEMENTS = {
    "Chamber-aware framing and deployed NCP_v3": APPROVED_TITLE,
    "Figure 1 — NCP_v3 construction, healthy baseline, and HCA carriers": (
        f"Figure 1 — {APPROVED_TITLE}"
    ),
    "Fig 1. Normal Chamber Prior v3 (NCP_v3):": (
        f"Fig 1. {APPROVED_TITLE}."
    ),
}


def migrate_text(text: str) -> str:
    if MIGRATION_NOTE_FRAGMENT in text:
        return text
    for prefix, replacement in SPECIAL_PREFIX_REPLACEMENTS.items():
        if text.startswith(prefix):
            if prefix.startswith("Fig 1."):
                return replacement
            text = replacement + text[len(prefix):]
            break
    text = text.replace(
        "Normal Chamber Prior, version 3 (NCP_v3)",
        "Normal Chamber Prior (NCP)",
    )
    text = text.replace(
        "Normal Chamber Prior v3 (NCP_v3)",
        "Normal Chamber Prior (NCP)",
    )
    text = STANDALONE_NCP_V3.sub("NCP", text)
    text = text.replace(
        "Fig. 1F HCA carrier prior",
        "predefined HCA carrier prior (Methods)",
    )
    text = text.replace(
        "the Fig. 1F HCA carrier prior",
        "the predefined HCA carrier prior (Methods)",
    )
    text = text.replace(
        "main-text Fig. 1E",
        "the prespecified selection procedure in Methods",
    )
    return text


def set_paragraph_text(paragraph, migrated: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(migrated)
        return
    paragraph.runs[0].text = migrated
    for run in paragraph.runs[1:]:
        run.text = ""


def migrate_paragraph(paragraph) -> int:
    original = paragraph.text
    migrated = migrate_text(original)
    if migrated == original:
        return 0
    # Run-level replacement preserves formatting for the common case.
    for run in paragraph.runs:
        run.text = migrate_text(run.text)
    if paragraph.text != migrated:
        set_paragraph_text(paragraph, migrated)
    return 1


def migrate_docx(path: Path) -> int:
    document = Document(path)
    changed = 0
    for paragraph in document.paragraphs:
        changed += migrate_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    changed += migrate_paragraph(paragraph)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                changed += migrate_paragraph(paragraph)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            changed += migrate_paragraph(paragraph)
    if changed:
        document.save(path)
    return changed


def main() -> None:
    roots = [
        ROOT / "manuscript_drafts",
        ROOT / "docs" / "manuscript" / "drafts",
        ROOT / "Manuscript_outputs",
    ]
    paths = sorted(
        {
            path
            for base in roots
            if base.is_dir()
            for path in base.rglob("*.docx")
        }
    )
    changed_files = 0
    changed_blocks = 0
    for path in paths:
        changed = migrate_docx(path)
        if changed:
            changed_files += 1
            changed_blocks += changed
            print(f"Migrated {path.relative_to(ROOT)} ({changed} text blocks)")
    print(
        f"Publication nomenclature migration complete: "
        f"{changed_files} files, {changed_blocks} text blocks"
    )


if __name__ == "__main__":
    main()
