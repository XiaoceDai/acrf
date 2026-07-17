#!/usr/bin/env python3
"""Supplementary Table  Sub-mechanism inventory underlying Fig 5E dual-axis"""
from __future__ import annotations

from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
OUT_TSV = ROOT / "results" / "n2d_sub_mechanism_table.tsv"
OUT_DOC = ROOT / "Manuscript_outputs" / \
            "Supplementary_Table_Sub_Mechanisms.docx"


SUB_MECHANISMS = [
    {
        "id": "SM-1",
        "axis": "RA  (acute, structural)",
        "name": "TGF-β / SLRP fibrotic invasion",
        "family": "F10  TGF-β  +  GAG metabolism family",
        "carrier": "fibroblast",
        "bulk":     "F10 pooled g = +0.68  CI [+0.11, +1.26]  ★  (Fig 2D)",
        "snrna":    "F10 × fibroblast  Δ = +0.010  p = 0.023  ★  (Fig 3B/C)",
        "spatial":  "GAG × fibroblast  ROI ρ = +0.43  p = 0.036  ★  (Fig 5D)",
        "duration": "non-significant (acute, not chronicity-driven)",
        "gene":     "DCN, PRELP among 24 reversal panel  (Fig 6B / 6C)",
        "test":     "N2d Test 2  SLRP enrichment  OR = 111.6  p = 2.4 × 10⁻⁴",
        "interpretation":
            "TGF-β cascade in LA fibroblasts induces small leucine-rich "
            "proteoglycans (SLRPs:  DCN, PRELP, OMD, FMOD class), driving "
            "ECM reorganization → matrix invasion of LA tissue.",
    },
    {
        "id": "SM-2",
        "axis": "RA  (acute, structural)",
        "name": "BMP signalling  ·  bidirectional",
        "family": "F11  BMP signalling family",
        "carrier": "adipocyte (gain)  +  lymphatic endothelial (loss)",
        "bulk":     "F11 pooled g = +0.85  CI [+0.27, +1.43]  ★  (Fig 2D)",
        "snrna":    "F11 × adipocyte   Δ = +0.032  p = 0.043  ★  (Fig 3B/C)\n"
                     "F11 × endothelial Δ = −0.009  p = 0.649  (HCA carrier ref)",
        "spatial":  "BMP-related programs co-vary in same compartments  (Fig 5D)",
        "duration": "non-significant",
        "gene":     "n/a  (program-level only)",
        "test":     "N2d Test 3  Δ_adipo > Δ_lymph in 8/9 paired donors\n"
                     "Wilcoxon p = 0.008  ★",
        "interpretation":
            "BMP signalling exhibits cell-context-dependent direction in AF: "
            "gain in atrial adipose tissue (paracrine candidate) and loss in "
            "lymphatic endothelium (atrial lymphatic dysfunction).",
    },
    {
        "id": "SM-3",
        "axis": "RA  (acute, structural)",
        "name": "HS-GAG metabolism  ·  diffuse",
        "family": "F9  HS-GAG / RUNX2 / B/2-secretin family",
        "carrier": "diffuse  (no single LA cell type passes p < 0.05)",
        "bulk":     "F9 pooled g = +0.79  CI [+0.22, +1.37]  ★  (Fig 2D)",
        "snrna":    "all cell types Δ ≈ 0  (Fig 3B / 5C BULK-LED class)",
        "spatial":  "n/a  (no cell-specific ROI signal)",
        "duration": "non-significant",
        "gene":     "n/a",
        "test":     "Fig 5C class assignment  =  BULK-LED",
        "interpretation":
            "Heparan-sulfate proteoglycan metabolism is broadly perturbed "
            "across the LA tissue without a single dominant cell carrier.",
    },
    {
        "id": "SM-4",
        "axis": "LA  (chronic, metabolic)",
        "name": "Mito Ca²⁺ buffering decay",
        "family": "F2  Mitochondrial Ca²⁺ transport",
        "carrier": "atrialCM  (diffuse, time-integrated)",
        "bulk":     "F2 pooled g = +0.11  CI crosses 0  (Fig 2D)",
        "snrna":    "Fig 3B no acute carrier",
        "spatial":  "Fig 5D atrialCM-correlated programs trend negative",
        "duration": "F2 LAA  ρ = -0.56  p = 0.06  (Fig 4B)\n"
                     "Fig 5C class:  DURATION-LED  (only LA family classified)",
        "gene":     "n/a",
        "test":     "Fig 4B Spearman  +  Fig 5C duration tier",
        "interpretation":
            "Mitochondrial Ca²⁺ buffering capacity declines monotonically "
            "with AF duration in LAA atrialCM, predisposing cytosolic "
            "Ca²⁺ overload and arrhythmogenic substrate maintenance.",
    },
    {
        "id": "SM-5",
        "axis": "LA  (chronic, metabolic)",
        "name": "Bioenergetic / OXPHOS decline",
        "family": "F7  TCA cycle / OXPHOS  +  F8  Complex I biogenesis",
        "carrier": "atrialCM",
        "bulk":     "F7  pooled g = -0.12  ;  F8 pooled g = -0.18  (sub-thresh)",
        "snrna":    "Fig 3B no acute carrier",
        "spatial":  "Fig 5D OXPHOS × atrialCM  ROI ρ = +0.45  p = 0.028  ★",
        "duration": "F7 LAA  ρ = -0.58  p = 0.06  (Fig 4C scatter)",
        "gene":     "n/a",
        "test":     "Fig 5D  Spearman  ρ  in cardiac AA ROIs",
        "interpretation":
            "Generic mitochondrial bioenergetics (TCA + OXPHOS + Complex I) "
            "trend downward with chronicity in LAA, spatially concentrated "
            "in atrialCM compartments.",
    },
    {
        "id": "SM-6",
        "axis": "LA  (small, consistent)",
        "name": "Sarcomeric loss  (LA → RA)",
        "family": "n/a  (gene-level)",
        "carrier": "atrialCM (likely)",
        "bulk":     "ACTA1 aCAI healthy = +0.44   AF flips to mean -0.24",
        "snrna":    "n/a",
        "spatial":  "n/a",
        "duration": "n/a",
        "gene":     "ACTA1  in  3 LA→RA reversal genes  (Fig 6B)",
        "test":     "Fig 6B chamber-reversal call",
        "interpretation":
            "Sarcomeric actin (ACTA1) loses LA bias in AF → contractile "
            "machinery dedifferentiation in LA atrium.",
    },
    {
        "id": "SM-7",
        "axis": "LA  (small, consistent)",
        "name": "Imprinting / adhesion loss  (LA → RA)",
        "family": "n/a",
        "carrier": "n/a (mixed)",
        "bulk":     "ZFP57  +  PCDH20  aCAI flip  in  2/3 cohorts",
        "snrna":    "n/a",
        "spatial":  "n/a",
        "duration": "n/a",
        "gene":     "ZFP57, PCDH20  in  3 LA→RA reversal genes",
        "test":     "Fig 6B chamber-reversal call",
        "interpretation":
            "Imprinting (ZFP57) and protocadherin adhesion (PCDH20) lose "
            "LA-bias in AF → epigenetic and junctional dedifferentiation.",
    },
    {
        "id": "SM-CROSS",
        "axis": "Cross-axis  (directional asymmetry)",
        "name": "RA → LA  matrix-invasion bias  (vs LA → RA  loss)",
        "family": "all  (gene-level)",
        "carrier": "fibroblast-dominant",
        "bulk":     "21 RA→LA reversal genes  vs  3 LA→RA (24 total)",
        "snrna":    "fibroblast top carrier  mean Δ_aligned = +0.042  (Fig 7d)",
        "spatial":  "consistent with Fig 5D fibroblast-program ρ pattern",
        "duration": "n/a",
        "gene":     "Fig 7a–c 24-gene panel  (binomtest p = 2.8 × 10⁻⁴)",
        "test":     "N2d Test 1  binomtest  21:3 vs 1:1  p = 2.8 × 10⁻⁴  ★",
        "interpretation":
            "AF rewires LA tissue toward an RA-like matrix-rich phenotype, "
            "not the reverse;  the 7:1 RA→LA : LA→RA ratio is far above "
            "chance and identifies AF chamber-reversal as a directional "
            "(not symmetric) process.",
    },
]


def main() -> None:
    df = pd.DataFrame(SUB_MECHANISMS)
    OUT_TSV.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT_TSV, sep="\t", index=False)
    print(f"wrote {OUT_TSV.relative_to(ROOT)}")

    # Optional: build a .docx supp table if python-docx is available.
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx not installed; skipping .docx output")
        return

    OUT_DOC.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.left_margin  = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.top_margin   = Cm(1.2)
    sec.bottom_margin = Cm(1.2)

    h = doc.add_heading(
        "Supplementary Table  Sub-mechanism inventory underlying "
        "Fig 5E dual-axis mechanism map", level=1)
    for run in h.runs:
        run.font.size = Pt(11)
    p = doc.add_paragraph(
        "Each sub-mechanism (SM-1 to SM-7 + cross-axis SM-CROSS) lists "
        "the supporting evidence layer-by-layer, including the per-test "
        "p-value from the N2d statistics module.  ★ marks p < 0.05.")
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(8)

    cols = ["id", "axis", "name", "family", "carrier",
              "bulk", "snrna", "spatial", "duration", "gene",
              "test", "interpretation"]
    headers = ["SM", "Axis", "Sub-mechanism", "Family / Set",
                 "Cell carrier", "Bulk (Fig 2)", "snRNA (Fig 3)",
                 "Spatial (Fig 5D)", "Duration (Fig 4)",
                 "Gene-level (Fig 6)", "Statistical test (N2d)",
                 "Biological interpretation"]
    table = doc.add_table(rows=1 + len(df), cols=len(cols))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for j, h_lab in enumerate(headers):
        hdr[j].text = h_lab
        for run in hdr[j].paragraphs[0].runs:
            run.font.size = Pt(7)
            run.bold = True
    for i, row in df.iterrows():
        cells = table.rows[i + 1].cells
        for j, c in enumerate(cols):
            cells[j].text = str(row[c]) if pd.notna(row[c]) else ""
            for run in cells[j].paragraphs[0].runs:
                run.font.size = Pt(6.5)
    doc.save(OUT_DOC)
    print(f"wrote {OUT_DOC.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
