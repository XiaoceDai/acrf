#!/usr/bin/env python3
"""Audit publication assets for visible historical NCP version suffixes."""
from __future__ import annotations

import hashlib
import re
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
PACK = ROOT / "submission" / "PLOS_CB_submission_pack"
OUT_TSV = RESULTS / "publication_ncp_nomenclature_qc.tsv"
OUT_MD = RESULTS / "publication_ncp_nomenclature_qc.md"

MIGRATION_NOTE = (
    "Publication-facing text uses NCP. Internal artifact names retain the "
    "historical ncp_v3 identifier to preserve frozen lineage, checksums and "
    "backward-compatible workflow paths."
)
EXPECTED_FIGURE1_TITLE = (
    "Definition and freezing of the healthy paired-chamber NCP coordinate system"
)
FORBIDDEN = re.compile(r"\bNCP(?:_|\s)V?3\b", flags=re.IGNORECASE)
INTERNAL_SUFFIXES = (
    ".tsv",
    ".csv",
    ".py",
    ".r",
    ".rds",
    ".md",
    ".txt",
    ".json",
    ".yaml",
    ".yml",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def docx_text(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.extend(cell.text for cell in row.cells)
    return "\n".join(blocks)


def pdf_text(path: Path) -> str:
    completed = subprocess.run(
        ["pdftotext", str(path), "-"],
        check=True,
        capture_output=True,
        text=True,
    )
    return completed.stdout


def token_around(text: str, start: int, end: int) -> str:
    left = start
    right = end
    while left > 0 and not text[left - 1].isspace():
        left -= 1
    while right < len(text) and not text[right].isspace():
        right += 1
    return text[left:right].strip("`'\"()[]{};,")


def visible_violations(text: str) -> list[str]:
    violations = []
    for match in FORBIDDEN.finditer(text):
        line_start = text.rfind("\n", 0, match.start()) + 1
        line_end = text.find("\n", match.end())
        if line_end < 0:
            line_end = len(text)
        line = text[line_start:line_end].strip()
        if MIGRATION_NOTE in text and "historical ncp_v3 identifier" in line:
            continue
        token = token_around(text, match.start(), match.end())
        token_lower = token.lower()
        if "/" in token or token_lower.endswith(INTERNAL_SUFFIXES):
            continue
        # Longer underscore-delimited names are internal object identifiers.
        if token_lower not in {"ncp_v3", "ncp v3"}:
            continue
        violations.append(line or token)
    return violations


def publication_assets(root: Path) -> list[tuple[Path, str]]:
    candidates: list[tuple[Path, str]] = [
        (root / "figures" / "fig1_legend.txt", "text"),
        (root / "results" / "figure1_algorithm_panel_definition.tsv", "text"),
        (root / "results" / "figure1_publication_asset_map.tsv", "text"),
        (root / "results" / "reproducibility_table.tsv", "text"),
        (root / "manuscript_drafts" / "Figure_Captions.docx", "docx"),
        (root / "manuscript_drafts" / "Data_Code_Availability.docx", "docx"),
        (root / "figures" / "fig1a_donor_paired_estimand.svg", "text"),
        (
            root
            / "figures"
            / "ncp_v3_singscore_reanalysis"
            / "fig1b_healthy_reference_direction_magnitude.svg",
            "text",
        ),
        (root / "figures" / "fig1c_selection_provenance.svg", "text"),
        (root / "figures" / "fig1d_program_family_topology.svg", "text"),
        (root / "figures" / "fig1a_donor_paired_estimand.pdf", "pdf"),
        (
            root
            / "figures"
            / "ncp_v3_singscore_reanalysis"
            / "fig1b_healthy_reference_direction_magnitude.pdf",
            "pdf",
        ),
        (root / "figures" / "fig1c_selection_provenance.pdf", "pdf"),
        (root / "figures" / "fig1d_program_family_topology.pdf", "pdf"),
    ]
    asset_map_path = root / "results" / "figure1_publication_asset_map.tsv"
    if asset_map_path.is_file():
        asset_map = pd.read_csv(asset_map_path, sep="\t")
        for column in ("pdf", "svg"):
            candidates.extend(
                (root / relative_path, column)
                for relative_path in asset_map[column].astype(str)
            )
    if PACK.is_dir():
        candidates.extend(
            (path, "docx") for path in sorted(PACK.rglob("*.docx"))
        )
        candidates.extend(
            (path, "pdf")
            for path in sorted((PACK / "02_Main_figures").glob("Figure_1*.pdf"))
        )
        candidates.append((PACK / "01_Manuscript_Word" / "fig1_legend.txt", "text"))
        candidates.append((PACK / "00_README_SUBMISSION_INDEX.txt", "text"))
        candidates.extend(
            (path, "text")
            for path in sorted((PACK / "05_Key_results_TSV").glob("*.tsv"))
        )
    unique: dict[Path, str] = {}
    for path, kind in candidates:
        unique[path] = kind
    return sorted(unique.items(), key=lambda item: str(item[0]))


def verify_frozen_lineage(root: Path) -> dict[str, str]:
    panel_path = root / "results" / "innovation_ncp_v3_panel.tsv"
    family_path = root / "results" / "innovation_ncp_v3_axis_families.tsv"
    fig1c_path = root / "results" / "fig1c_selection_provenance_source.tsv"
    fig1d_path = root / "results" / "fig1d_program_family_topology_source.tsv"
    for path in (panel_path, family_path, fig1c_path, fig1d_path):
        if not path.is_file():
            raise AssertionError(f"Required frozen lineage path is missing: {path}")

    panel = pd.read_csv(panel_path, sep="\t")
    family = pd.read_csv(family_path, sep="\t")
    fig1c = pd.read_csv(fig1c_path, sep="\t")
    fig1d = pd.read_csv(fig1d_path, sep="\t")
    frozen = set(panel["pathway"].astype(str))
    if len(frozen) != 30:
        raise AssertionError("Frozen panel no longer contains 30 unique programs")
    if set(fig1c.loc[fig1c["final_selected"], "program_id"].astype(str)) != frozen:
        raise AssertionError("Figure 1C no longer resolves to the frozen panel")
    if set(fig1d["program_id"].astype(str)) != frozen:
        raise AssertionError("Figure 1D no longer resolves to the frozen panel")
    if set(family["pathway"].astype(str)) != frozen or family["family_id"].nunique() != 16:
        raise AssertionError("Figure 1D family mapping no longer resolves 30 programs to 16 families")
    return {
        str(panel_path.relative_to(root)): sha256(panel_path),
        str(family_path.relative_to(root)): sha256(family_path),
    }


def verify_labels(root: Path) -> None:
    labels_path = (
        root
        / "results"
        / "ncp_v3_singscore_reanalysis"
        / "fig1b_canonical_display_order.tsv"
    )
    labels = set(
        pd.read_csv(labels_path, sep="\t")["publication_label"].astype(str)
    )
    expected = {
        "Mito Ca²⁺ transport",
        "Mito β-oxidation",
        "TGF-β signalling (Hall)",
        "NTRK2 / TrkB signalling",
    }
    missing = expected - labels
    if missing:
        raise AssertionError(f"Publication labels differ from Figure 1B: {sorted(missing)}")


def run_qc(root: Path = ROOT) -> pd.DataFrame:
    rows = []
    all_violations: list[str] = []
    for path, kind in publication_assets(root):
        if not path.is_file():
            raise AssertionError(f"Publication asset is missing: {path}")
        if kind == "docx":
            text = docx_text(path)
        elif kind == "pdf":
            text = pdf_text(path)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
        violations = visible_violations(text)
        rows.append(
            {
                "artifact_path": str(path.relative_to(root)),
                "artifact_kind": kind,
                "sha256": sha256(path),
                "visible_NCP_v3_occurrences": len(violations),
                "status": "PASS" if not violations else "FAIL",
            }
        )
        all_violations.extend(
            f"{path.relative_to(root)}: {violation}" for violation in violations
        )

    legend = (root / "figures" / "fig1_legend.txt").read_text(encoding="utf-8")
    if EXPECTED_FIGURE1_TITLE not in legend:
        raise AssertionError("Figure 1 overall title is not the approved title")
    verify_labels(root)
    frozen_hashes = verify_frozen_lineage(root)
    readme = (root / "README.md").read_text(encoding="utf-8")
    if MIGRATION_NOTE not in readme.replace("\n", " "):
        raise AssertionError("README is missing the approved nomenclature migration note")
    asset_map = pd.read_csv(
        root / "results" / "figure1_publication_asset_map.tsv", sep="\t"
    )
    if asset_map["panel"].tolist() != [
        "Figure 1 composite",
        "Figure 1A",
        "Figure 1B",
        "Figure 1C",
        "Figure 1D",
    ]:
        raise AssertionError("Figure 1 publication asset map is not composite plus A-D provenance")
    for column in ("pdf", "svg", "png"):
        for relative_path in asset_map[column]:
            if not (root / relative_path).is_file():
                raise AssertionError(f"Figure 1 mapped asset is missing: {relative_path}")
    if all_violations:
        raise AssertionError(
            "Visible historical NCP nomenclature detected:\n"
            + "\n".join(all_violations)
        )

    report = pd.DataFrame(rows)
    report.to_csv(OUT_TSV, sep="\t", index=False)
    hash_lines = "\n".join(
        f"- `{path}` — `{digest}`" for path, digest in frozen_hashes.items()
    )
    OUT_MD.write_text(
        "# Publication-facing NCP nomenclature QC\n\n"
        "Status: **PASS**\n\n"
        f"- publication assets scanned: {len(report)}\n"
        "- visible `NCP_v3`, `NCP V3`, or `NCP_V3` occurrences: 0\n"
        "- internal paths, filenames, manifests, and object IDs: permitted\n"
        "- Figure 1 overall title: approved\n"
        "- Figure 1B/1D mathematical and spelling labels: exact\n"
        "- Figure 1C frozen program set: unchanged\n"
        "- Figure 1D 30-program/16-family topology: unchanged\n\n"
        "## Frozen internal source hashes\n"
        f"{hash_lines}\n",
        encoding="utf-8",
    )
    return report


def main() -> None:
    report = run_qc()
    print(
        f"Publication NCP nomenclature QC: PASS "
        f"({len(report)} assets; 0 visible historical suffixes)"
    )


if __name__ == "__main__":
    main()
