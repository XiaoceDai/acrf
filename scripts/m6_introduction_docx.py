#!/usr/bin/env python3
"""
Build manuscript_drafts/Introduction.docx

Phase 7 locked Introduction: five paragraphs; citation set [1-3], [4,5],
[6-10], [11-13], [14] retained without bibliography renumbering.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from m0_doc_style import apply_base_style

OUT = ROOT / "manuscript_drafts" / "Introduction.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


PARAGRAPHS: list[str] = [
    (
        "Atrial fibrillation (AF) is the most common sustained cardiac arrhythmia and a "
        "major contributor to cardiovascular morbidity [1-3]. Its substrate can reflect "
        "different combinations of electrical instability, structural remodeling, "
        "inflammation, and metabolic stress [4,5]. Human transcriptomic studies have "
        "repeatedly implicated extracellular-matrix organization, ion-channel and "
        "calcium-handling pathways, immune signaling, and mitochondrial, metabolic, or "
        "translational programs [6-10]. However, comparison across cohorts remains "
        "difficult because studies differ in tissue sampling, chamber composition, "
        "clinical context, molecular platform, and analytical design."
    ),
    (
        "One persistent source of heterogeneity is atrial chamber identity. The left and "
        "right atria have distinct baseline transcriptional programs and can show "
        "chamber-specific responses in AF [11-13]. Pooling chambers, comparing cohorts "
        "with different chamber compositions, or treating paired left- and right-atrial "
        "samples as independent observations can therefore confound disease-associated "
        "change with healthy LA–RA asymmetry. In paired atrial studies, rhythm is assigned "
        "at donor level whereas chamber is measured within donor, making donor-level "
        "LA–RA contrasts a natural unit for chamber-aware disease comparisons."
    ),
    (
        "Pathway- and program-level summaries can reduce transcriptomic dimensionality, "
        "but raw chamber contrasts are not directly comparable across programs with "
        "different healthy baseline magnitudes. A reusable framework should therefore "
        "define each program’s chamber direction in a healthy paired reference, scale "
        "disease-associated asymmetry by a program-specific healthy magnitude, and freeze "
        "program membership before disease-cohort analysis. Applying the same definitions "
        "across cohorts avoids redefining the molecular coordinate system in response to "
        "each dataset."
    ),
    (
        "Cross-dataset integration also requires preserving the distinctions among "
        "evidence layers. Paired bulk tissue, single-nucleus profiles, cross-sectional "
        "duration data, spatial regions of interest, and genetic summary statistics differ "
        "in analysis unit, estimand, and uncertainty structure. Collapsing them into a "
        "single score can obscure the difference between reproducible effect-size "
        "patterns, nominal cellular localization, descriptive chronicity trends, "
        "ROI-level spatial context, and genetic calibration. A useful integrative "
        "framework should make cross-layer agreement visible while retaining "
        "layer-specific interpretation boundaries."
    ),
    (
        "Here we developed the Normal Chamber Prior (NCP), a frozen healthy reference of "
        "30 LA–RA programs organized into 16 families from paired non-failing human atria "
        "[14]. We applied the same definitions to three paired-chamber bulk cohorts, LA "
        "single-nucleus RNA sequencing, a duration-annotated left-atrial-appendage cohort, "
        "GeoMx spatial ROIs, AF GWAS summary statistics, and atrial cis-eQTL resources. "
        "We asked which families showed reproducible bulk remodeling, where selected "
        "signals localized within LA, how family activity related to cross-sectional "
        "duration, and which families were prioritized by a prespecified three-layer rule "
        "(Path A). The analysis identified selective remodeling across both LA-biased "
        "translation and ribosome-quality-control programs and RA-biased matrix, "
        "proteoglycan, and developmental-signaling programs, yielded nominal LA "
        "compartment localization for four families, and prioritized two families across "
        "the three transcriptomic evidence layers. Spatial and genetic analyses supplied "
        "contextual and calibration-limited, non-nested traceability and did not enter "
        "Path A. The NCP therefore provides a reusable chamber-aware coordinate system for "
        "comparing heterogeneous AF datasets without redefining molecular programs in each "
        "cohort."
    ),
]


def main() -> None:
    doc = Document()
    apply_base_style(doc)

    doc.add_heading("Introduction", level=1)

    for p in PARAGRAPHS:
        para = doc.add_paragraph(p)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
