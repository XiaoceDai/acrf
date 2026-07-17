#!/usr/bin/env python3
"""Build manuscript_drafts/Supplementary_Table_S6_anchor_matrix.docx from MR/coloc TSV."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm

from m0_doc_style import apply_base_style, add_body, add_table

IN = ROOT / "results" / "figure6_mr_anchors_v3_pipeline_wide.tsv"
OUT = ROOT / "manuscript_drafts" / "Supplementary_Table_S6_anchor_matrix.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


def _fmt_cell(v) -> str:
    if pd.isna(v):
        return ""
    if isinstance(v, float):
        if abs(v) >= 1e4 or (abs(v) > 0 and abs(v) < 1e-4):
            return f"{v:.6e}"
        return f"{v:.6g}"
    return str(v)


def main() -> None:
    df = pd.read_csv(IN, sep="\t")
    doc = Document()
    apply_base_style(doc)

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_height, sec.page_width = sec.page_width, sec.page_height
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)

    doc.add_heading(
        "Supplementary Table S6 | cis-eQTL × AF-GWAS colocalisation anchor matrix (v3)",
        level=1,
    )
    add_body(
        doc,
        "Full numeric output for pre-registered framework anchors and calibration "
        "genes, primary tissue (GTEx atrial appendage; eQTL Catalogue QTD000251) "
        "and blood sensitivity (QTD000356), including coloc posterior "
        "probabilities and pipeline verdict fields. Source file: "
        "results/figure6_mr_anchors_v3_pipeline_wide.tsv. "
        "Manuscript-facing slim matrix: results/figure6_mr_anchors_v3_consolidated.tsv.",
    )

    headers = [str(c) for c in df.columns]
    rows: list[list[str]] = [headers]
    for _, r in df.iterrows():
        rows.append([_fmt_cell(r[c]) for c in df.columns])

    # Wide matrix: no fixed column widths; use landscape + Word auto layout.
    add_table(doc, rows)

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({len(df)} rows × {len(df.columns)} cols)")


if __name__ == "__main__":
    main()
