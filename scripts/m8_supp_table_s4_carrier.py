#!/usr/bin/env python3
"""Build Supplementary Table S4 — per-gene per-compartment AF − Ctrl
expression Δ for the 24 chamber-reversal genes (Fig. 6d source data)."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import pandas as pd
from docx import Document

from m0_doc_style import apply_base_style, add_body, add_table

OUT = ROOT / "manuscript_drafts" / "Supplementary_Table_S4_reversal_carrier.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)
IN_LONG = ROOT / "results" / "n2b_reversal_per_celltype_long.tsv"
IN_GENES = ROOT / "results" / "n2_chamber_reversal_genes.tsv"


def main() -> None:
    long = pd.read_csv(IN_LONG, sep="\t")
    rev = pd.read_csv(IN_GENES, sep="\t")
    rev_dir = rev.set_index("gene")["healthy_direction"]

    pivot = (long.dropna(subset=["delta_expr"])
                  .pivot_table(index="gene",
                               columns="compartment",
                               values="delta_expr"))
    cols = ["atrialCM", "fibroblast",
            "endothelial_or_vascular", "immune"]
    pivot = pivot[cols]
    pivot["healthy_direction"] = rev_dir.reindex(pivot.index)
    pivot = pivot.sort_values(
        ["healthy_direction", "fibroblast"],
        ascending=[True, False],
    )

    doc = Document()
    apply_base_style(doc)
    doc.add_heading(
        "Supplementary Table S4 | Per-gene per-compartment AF − Ctrl "
        "expression Δ for the 24 chamber-reversal genes",
        level=1,
    )
    add_body(
        doc,
        "Source data for Fig. 6d. Per gene, per compartment Δ_expr = "
        "mean(expression in AF biosamples) − mean(expression in Ctrl "
        "biosamples) on the GSE255612 single-nucleus pseudobulks "
        "(AF n = 18, Ctrl n = 16; expression units = pseudobulk-mean "
        "log-normalised counts as released by the project). Per "
        "compartment summary: mean across the 21 RA → LA matrix-"
        "invasion genes — fibroblast +0.045, immune +0.022, atrialCM "
        "+0.006, endothelial / vascular +0.006 — identifies fibroblast "
        "as the carrier compartment of the matrix-invasion signature. "
        "Healthy direction is inherited from Supplementary Table S3.",
    )

    rows = [["gene", "healthy direction",
             "Δ atrialCM", "Δ fibroblast",
             "Δ endothelial / vasc", "Δ immune"]]
    for g, r in pivot.iterrows():
        rows.append([
            g, str(r["healthy_direction"]),
            f"{r['atrialCM']:+.4f}",
            f"{r['fibroblast']:+.4f}",
            f"{r['endothelial_or_vascular']:+.4f}",
            f"{r['immune']:+.4f}",
        ])
    add_table(doc, rows,
              col_widths_cm=[2.0, 2.4, 2.4, 2.6, 3.0, 2.0])

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({len(pivot)} rows)")


if __name__ == "__main__":
    main()
