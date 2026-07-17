#!/usr/bin/env python3
"""Build manuscript_drafts/Supplementary_Table_S1_cohort_demographics.docx from the cohort TSV.

Submission workflow (PLOS / similar):
  1. Run: python scripts/p1_demographics_table.py   → results/cohort_demographics_table.tsv
  2. Run: python scripts/m_supp_table_s1_cohort_demographics.py   → this .docx
  3. In Word/LibreOffice: export PDF if the journal requires a PDF supporting file.

The TSV remains the versioned source of truth in the repository; the DOCX/PDF are
uploaded as Supporting Information Table S1.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import pandas as pd
from docx import Document

from m0_doc_style import apply_base_style, add_body, add_table

IN = ROOT / "results" / "cohort_demographics_table.tsv"
OUT = ROOT / "manuscript_drafts" / "Supplementary_Table_S1_cohort_demographics.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


def main() -> None:
    df = pd.read_csv(IN, sep="\t")
    doc = Document()
    apply_base_style(doc)

    doc.add_heading(
        "Supplementary Table S1 | Public cohorts and GEO accessions",
        level=1,
    )
    add_body(
        doc,
        "Cohort-level summary for all publicly available RNA-related resources "
        "used in the manuscript (healthy reference, bulk discovery/replication, "
        "single-nucleus, multiome gate, and spatial ROI cohorts). Accession "
        "identifiers are listed in the first column. For GSE128188, GSE79768, "
        "and GSE115574, donor counts and the per-donor companion file "
        "(results/cohort_demographics_long.tsv) are intersected with "
        "results/innovation_ncp_v3_acai_long.tsv when tables are regenerated, "
        "so Fig. 2, Supplementary Fig. S1, and this table agree on analytical "
        "sample sizes (GSE115574: 13 AF + 15 SR after paired-chamber / coverage "
        "QC; GEO registry lists additional donors). Regenerate from "
        "results/cohort_demographics_table.tsv after editing "
        "scripts/p1_demographics_table.py.",
    )

    headers = list(df.columns)
    rows: list[list[str]] = [headers]
    for _, r in df.iterrows():
        rows.append(["" if pd.isna(r[c]) else str(r[c]) for c in headers])

    add_table(
        doc,
        rows,
        col_widths_cm=[2.0, 3.2, 3.0, 2.6, 1.0, 1.0, 1.0, 2.2, 1.6, 4.5],
    )

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({len(df)} cohort rows)")


if __name__ == "__main__":
    main()
