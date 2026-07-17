#!/usr/bin/env python3
"""
Build manuscript_drafts/Data_Code_Availability.docx

Phase 9B: Option B Data / Code Availability only.
No Zenodo DOI, provisional commit, license claim, or unverified release identifier.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from m0_doc_style import apply_base_style, add_body

OUT = ROOT / "manuscript_drafts" / "Data_Code_Availability.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


DATA_AVAILABILITY = (
    "All analyzed molecular and genetic resources in this study are publicly "
    "available from their original repositories. Dataset accessions, providers, "
    "and analysis roles are listed in S1 Table. No new participants or "
    "biospecimens were generated for this study. Public raw datasets remain "
    "hosted by the original repositories and are not redistributed here. After "
    "release freeze, the project release will provide processed Normal Chamber "
    "Prior (NCP) framework outputs, figure-level source tables, frozen metadata, "
    "formal Supporting Information workbooks (S1–S6 Table), provenance records "
    "(including the GCST90624411 GWAS provenance package and the external AF "
    "gene-set manifest), and reproducibility materials. Full GWAS summary "
    "statistics should be obtained from the GWAS Catalog accession GCST90624411 "
    "rather than treated as a manuscript supplement. Dataset-specific analysis "
    "units are retained; these resources do not constitute one pooled cohort. "
    "No identifiable participant-level dataset is distributed."
)

CODE_AVAILABILITY = (
    "Analysis code is distributed as the installable Python package acrf "
    "(import name crf) at https://github.com/XiaoceDai/acrf. The project "
    "release will include analysis scripts together with figure-, table-, and "
    "manuscript-generation workflows, QC reports, and SHA256 manifests for "
    "frozen outputs. Final archival DOI, release commit, and release tag will "
    "be added only after release freeze and are not reported as provisional "
    "values here. The exact historical production coloc version could not be "
    "recovered from local metadata; the final release environment will record "
    "the coloc and Python dependency versions used for reproduction and must "
    "not be interpreted as proof of the historical production environment. "
    "Output SHA256 hashes will accompany the release for verification."
)


def main() -> None:
    doc = Document()
    apply_base_style(doc)

    doc.add_heading("Data and Code Availability", level=1)

    doc.add_heading("Data Availability", level=2)
    add_body(doc, DATA_AVAILABILITY)

    doc.add_heading("Code Availability", level=2)
    add_body(doc, CODE_AVAILABILITY)

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
