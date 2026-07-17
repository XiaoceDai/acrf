#!/usr/bin/env python3
"""
Build manuscript_drafts/Discussion.docx

Phase 8 locked Discussion: seven paragraphs; citation sets [5,6,9] and
[11-13] retained without bibliography renumbering.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from m0_doc_style import apply_base_style

OUT = ROOT / "manuscript_drafts" / "Discussion.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


PARAGRAPHS: list[str] = [
    (
        "This study establishes a chamber-aware coordinate system for comparing atrial "
        "fibrillation (AF) remodeling across heterogeneous human datasets. By freezing a "
        "healthy LA–RA reference before disease-cohort analysis, the Normal Chamber Prior "
        "(NCP) separated baseline chamber identity from disease-associated displacement "
        "of that identity. Cross-cohort reproducibility in paired bulk tissue was selective "
        "and spanned both sides of the healthy reference: LA-biased translation and "
        "ribosome-quality-control families and RA-biased matrix, proteoglycan, and "
        "developmental-signaling families were represented among the High-consistency "
        "group. LA single-nucleus analysis supplied nominal compartment context for four "
        "families, cross-sectional duration analysis supplied a separate chronicity layer, "
        "and the prespecified Path A rule prioritized F10 GAG disease / TGF-β and "
        "F15 NTRK2 / TrkB across the three transcriptomic evidence layers. GeoMx and "
        "genetic analyses added spatial and genetic context but did not enter Path A. "
        "Accordingly, the evidence supports a reproducible chamber-aware architecture "
        "rather than a single causal mechanism."
    ),
    (
        "Prior human AF studies have repeatedly implicated extracellular-matrix "
        "organization, inflammatory and signaling programs, and mitochondrial or "
        "metabolic remodeling [5,6,9]. The NCP addresses a complementary problem: these "
        "themes are difficult to compare when cohorts differ in chamber sampling or when "
        "healthy LA–RA asymmetry is incompletely modeled [11-13]. The aligned family scale "
        "also requires a specific interpretation. A negative pooled AF-minus-SR effect "
        "indicates displacement opposite to the frozen healthy chamber direction on the "
        "aligned scale; it does not imply uniform downregulation of every gene in the "
        "family. All six High-consistency pooled effects were negative on this scale, but "
        "the High set included both LA-biased translation/ribosome-quality-control "
        "families and RA-biased matrix/developmental families. The final bulk results "
        "therefore do not support a simple division in which structural families are "
        "uniformly reproducible and translational families are uniformly cohort dependent. "
        "Instead, they suggest that AF perturbs selected components of both healthy "
        "chamber-biased arms, while other mitochondrial and metabolic families remain more "
        "variable across cohorts."
    ),
    (
        "The LA single-nucleus layer showed that family-level bulk remodeling does not map "
        "uniformly to cell compartments. F10 retained its healthy fibroblast context and "
        "was classified as Fidelity; F11 localized to adipocytes and F15 to endothelial "
        "cells only, producing Relay patterns; and F16 involved fibroblasts together with "
        "atrial cardiomyocytes, producing Extension. These assignments were based on "
        "nominal donor-label permutation support in LA and were not FDR-controlled "
        "discoveries. They should therefore be interpreted as cellular localization cues, "
        "not as validated cell-type mechanisms or an independent replication of "
        "bi-chamber remodeling. Conversely, Below-detection indicates that no eligible LA "
        "compartment met the preset support rule; it does not imply absence of remodeling, "
        "particularly for families whose bulk reference is RA biased. The supplementary "
        "F4 BCAA example further shows that a constituent program can retain nominal "
        "compartment support while its parent family remains Below-detection, emphasizing "
        "that biological resolution depends on the level of aggregation."
    ),
    (
        "The duration and triangulation layers were similarly designed to preserve rather "
        "than erase differences in evidence type. Duration associations were estimated in "
        "12 AF-LAA samples and classified by the magnitude of the full-precision "
        "cross-sectional Spearman correlation. Only F2 and F9 met the Strong-trend "
        "threshold, and regional LAPV direction was not uniform, illustrating that trend "
        "magnitude and regional sign agreement are distinct properties. These observations "
        "do not establish progression, recurrence risk, treatment response, or temporal "
        "causal order. Path A then combined exactly three prespecified binary summaries: "
        "High bulk consistency, a non-Below-detection single-nucleus class, and a "
        "non-Stable duration class. F10 and F15 met all three conditions, whereas F9 met "
        "the bulk and duration conditions without LA single-nucleus localization. Path A "
        "is therefore best read as transparent rule-based triage across non-exchangeable "
        "layers, not as a joint statistical test or causal-validation criterion."
    ),
    (
        "GeoMx supplied an additional but deliberately separate spatial view. Nominal "
        "ROI-level summaries linked F10 with fibroblast-panel scores, F15 with "
        "macrophage-panel scores, and F9 with an AF-minus-SR difference. The F10 result is "
        "directionally compatible with its fibroblast single-nucleus localization, whereas "
        "the F15 macrophage-panel association and endothelial single-nucleus localization "
        "arise from different assays, signatures, and analysis units and should not be "
        "collapsed into a single cell-type claim. GeoMx family signatures also used an "
        "explicit family-gene-union scoring exception rather than the program-first family "
        "aggregation used in the primary transcriptomic analyses. Because the eligible "
        "data comprised multiple ROIs from each of four patients and no patient-clustered "
        "or mixed-effects inference was applied, these findings provide ROI-level spatial "
        "context only. They neither validate Path A at patient level nor constitute a "
        "fourth gate."
    ),
    (
        "The genetic analyses further illustrate why evidence layers should remain "
        "separate. Seven of 12 anchor windows contained local genome-wide AF associations, "
        "but only SMAD9 met the corrected atrial eGene criterion and no anchor reached the "
        "prespecified PP4 threshold for same-variant colocalization. All anchors were "
        "PP3-dominant, and lead AF-GWAS and atrial-eQTL variants were separated by "
        "hundreds of kilobases to nearly two megabases. These diagnostics are compatible "
        "with distinct local association signals but do not prove that different causal "
        "variants underlie the two traits. The failure of PITX2 and KCNN3 calibration "
        "controls to meet the same colocalization threshold underscores the conservatism "
        "of the available bulk-LAA regulatory setting. Similarly, the external gene-set "
        "analysis yielded only sparse CAV1 and HEY2 membership and no family-level "
        "enrichment after correction across 64 tests. Together, these results do not "
        "support same-variant bulk-LAA total-expression mediation for the tested anchors, "
        "but neither do they exclude cell-state-specific regulation, developmental "
        "effects, distal or trans mechanisms, rare variation, or other untested genetic "
        "routes. The non-nested traceability matrix was intended to expose these distinct "
        "forms of evidence without converting them into an ordinal tier, composite score, "
        "or genetics-based Path A gate."
    ),
    (
        "Several limitations define the appropriate scope of the framework. The healthy "
        "reference comprised eight donors, and frozen pathway definitions may not capture "
        "all cohort- or comorbidity-specific biology. Public bulk cohorts differed in "
        "platform and clinical context, and bulk chamber contrasts can reflect both "
        "within-cell regulation and shifts in cellular composition. The single-nucleus "
        "resource was LA only, used nominal localization thresholds, and could not test "
        "RA-driven compartment remodeling. Duration analysis was cross-sectional and "
        "based on 12 AF-LAA samples. GeoMx included four patients, multiple ROIs per "
        "patient, and no clustered inference; its reduced-panel structure also required a "
        "gene-union scoring exception. Genetic analyses used bulk atrial-appendage eQTL "
        "data rather than cell-state-specific regulatory maps, and exact historical "
        "software-environment recovery was incomplete for some dependencies. Finally, "
        "Path A depends on prespecified categorical boundaries and should be evaluated "
        "prospectively rather than treated as a clinical or mechanistic classifier. The "
        "principal contribution of the NCP is therefore methodological: frozen program "
        "definitions, donor-paired chamber alignment, program-specific healthy scaling, "
        "and explicit evidence-layer boundaries provide a reproducible framework for "
        "comparing AF datasets. Prospective paired-atrial studies with matched "
        "single-cell, spatial, regulatory, proteomic, metabolic, and perturbational data "
        "will be needed to test temporal order and causal mechanisms."
    ),
]


def main() -> None:
    if len(PARAGRAPHS) != 7:
        raise SystemExit(f"Expected exactly 7 Discussion paragraphs, got {len(PARAGRAPHS)}")

    doc = Document()
    apply_base_style(doc)

    doc.add_heading("Discussion", level=1)

    for p in PARAGRAPHS:
        para = doc.add_paragraph(p)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
